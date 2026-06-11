"""Tests du Lot 7.1 : ExportService — export complet des données d'un projet.

Couvre : l'export xlsx (classeur rouvert par openpyxl, TOUTES les feuilles
attendues présentes et non vides, sondes nombre de nœuds et valeur
d'évaluation contre la base), l'export csv (un ZIP avec un .csv par feuille,
utf-8-sig avec BOM, séparateur « ; »), la slugification du nom de projet
(accents translittérés, « / » retiré), le projet inconnu (ValueError) et la
création du répertoire de destination.
"""

from __future__ import annotations

import re
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from supplyscore.core.clock import FixedClock, iso_week
from supplyscore.domain.models import Project, TaskStatus
from supplyscore.services import SupplyScoreService
from supplyscore.services.decisions import DecisionService
from supplyscore.services.events import EventEngine
from supplyscore.services.exports import ExportService

#: Mercredi 2026-06-10 12:00 locale — horodatage de fichier « 20260610_120000 ».
_NOW = datetime(2026, 6, 10, 12, 0).timestamp()
_STAMP = "20260610_120000"

#: Feuilles attendues, dans l'ordre du classeur.
_FEUILLES = [
    "projet",
    "noeuds",
    "arcs",
    "jalons",
    "historique_urgences",
    "evaluations",
    "snapshots_kpi",
    "evenements",
    "decisions",
    "revues_hebdo",
    "audit",
    "audit_registre",
]


@pytest.fixture
def service(tmp_path: Path):
    svc = SupplyScoreService(db_dir=tmp_path / "store", clock=FixedClock(_NOW))
    yield svc
    svc.close()


@pytest.fixture
def projet(service: SupplyScoreService) -> Project:
    """Projet de démonstration reproductible (seed_demo)."""
    return service.seed_demo(n_ranks=2, seed=1)


@pytest.fixture
def node_id(service: SupplyScoreService, projet: Project) -> str:
    """Premier nœud actif (ordre déterministe) du projet de démo."""
    nodes = sorted(service.repo.nodes(), key=lambda n: n.id)
    actifs = [
        n for n in nodes if n.status is TaskStatus.ACTIVE and n.onboarding_state == "complete"
    ]
    assert actifs, "le projet de démo doit contenir au moins un nœud actif"
    return actifs[0].id


@pytest.fixture
def service_riche(service: SupplyScoreService, projet: Project, node_id: str) -> SupplyScoreService:
    """Service enrichi : événement, décision, revue hebdo et mutation registre.

    Garantit que les feuilles « evenements », « decisions », « revues_hebdo »,
    « snapshots_kpi », « audit » (bases client) et « audit_registre » ont au
    moins une ligne chacune.
    """
    EventEngine(service).apply(
        node_id, "pic_demande", {"nouvelle_demande": 120.0}, operator_id="op-1", notes="pic"
    )
    DecisionService(service).record(node_id, "Relancer le fournisseur", operator_id="op-2")
    semaine = iso_week(service.clock_for(projet.id).now())
    service.client_db(node_id).upsert_weekly_review(
        node_id, semaine, volets={"ahp": 1}, started_at=_NOW, completed_at=_NOW + 600.0
    )
    service.mutations.update_node_fields(
        node_id, {"location": "Zone-Test-Export"}, source="edit", operator_id="op-3"
    )
    return service


@pytest.fixture
def exports(service_riche: SupplyScoreService) -> ExportService:
    return ExportService(service_riche)


def _colonne(ws, header: str) -> int:
    """Index (1-based) de la colonne dont l'en-tête (ligne 1) vaut ``header``."""
    for cell in ws[1]:
        if cell.value == header:
            return cell.column
    raise AssertionError(f"en-tête introuvable dans {ws.title!r} : {header!r}")


# --- Export xlsx ----------------------------------------------------------------------


class TestExportXlsx:
    def test_fichier_cree_avec_le_bon_nom(
        self, exports: ExportService, projet: Project, tmp_path: Path
    ) -> None:
        path = exports.export_project(projet.id, fmt="xlsx", dest_dir=tmp_path / "exports")

        assert path.exists()
        assert re.fullmatch(rf"SupplyScore_[A-Za-z0-9-]+_{_STAMP}\.xlsx", path.name)

    def test_toutes_les_feuilles_presentes_et_non_vides(
        self, exports: ExportService, projet: Project, tmp_path: Path
    ) -> None:
        path = exports.export_project(projet.id, dest_dir=tmp_path / "exports")

        wb = load_workbook(path)
        assert wb.sheetnames == _FEUILLES
        for nom in _FEUILLES:
            assert wb[nom].max_row >= 2, f"feuille {nom!r} vide (en-têtes seuls)"

    def test_sonde_nombre_de_noeuds(
        self,
        exports: ExportService,
        service_riche: SupplyScoreService,
        projet: Project,
        tmp_path: Path,
    ) -> None:
        path = exports.export_project(projet.id, dest_dir=tmp_path / "exports")

        ws = load_workbook(path)["noeuds"]
        attendu = len(service_riche.registry.list_nodes(projet.id))
        assert ws.max_row - 1 == attendu

    def test_sonde_valeur_evaluation_egale_a_la_base(
        self,
        exports: ExportService,
        service_riche: SupplyScoreService,
        projet: Project,
        node_id: str,
        tmp_path: Path,
    ) -> None:
        node = service_riche.registry.get_node(node_id)
        assert node is not None
        derniere = service_riche.client_db(node_id).latest_assessment(node_id)
        assert derniere is not None

        path = exports.export_project(projet.id, dest_dir=tmp_path / "exports")

        ws = load_workbook(path)["evaluations"]
        col_noeud, col_ud = _colonne(ws, "nœud"), _colonne(ws, "ud")
        valeurs = [
            row[col_ud - 1].value
            for row in ws.iter_rows(min_row=2)
            if row[col_noeud - 1].value == node.name
        ]
        assert valeurs, "aucune ligne d'évaluation pour le nœud sondé"
        assert any(v == pytest.approx(derniere.ud) for v in valeurs)

    def test_sonde_evenement_et_decision_injectes(
        self,
        exports: ExportService,
        service_riche: SupplyScoreService,
        projet: Project,
        node_id: str,
        tmp_path: Path,
    ) -> None:
        node = service_riche.registry.get_node(node_id)
        assert node is not None

        path = exports.export_project(projet.id, dest_dir=tmp_path / "exports")
        wb = load_workbook(path)

        ws_ev = wb["evenements"]
        col_noeud, col_type = _colonne(ws_ev, "nœud"), _colonne(ws_ev, "type")
        types = [
            row[col_type - 1].value
            for row in ws_ev.iter_rows(min_row=2)
            if row[col_noeud - 1].value == node.name
        ]
        assert "pic_demande" in types

        ws_dec = wb["decisions"]
        col_desc = _colonne(ws_dec, "description")
        descriptions = [row[col_desc - 1].value for row in ws_dec.iter_rows(min_row=2)]
        assert "Relancer le fournisseur" in descriptions


# --- Export csv (zip) -------------------------------------------------------------------


class TestExportCsv:
    def test_zip_avec_un_csv_par_feuille(
        self, exports: ExportService, projet: Project, tmp_path: Path
    ) -> None:
        path = exports.export_project(projet.id, fmt="csv", dest_dir=tmp_path / "exports")

        assert path.suffix == ".zip"
        assert re.fullmatch(rf"SupplyScore_[A-Za-z0-9-]+_{_STAMP}\.zip", path.name)
        with zipfile.ZipFile(path) as archive:
            assert sorted(archive.namelist()) == sorted(f"{nom}.csv" for nom in _FEUILLES)

    def test_encodage_utf8_sig_et_separateur_point_virgule(
        self,
        exports: ExportService,
        service_riche: SupplyScoreService,
        projet: Project,
        tmp_path: Path,
    ) -> None:
        path = exports.export_project(projet.id, fmt="csv", dest_dir=tmp_path / "exports")

        with zipfile.ZipFile(path) as archive:
            brut = archive.read("noeuds.csv")
        assert brut.startswith(b"\xef\xbb\xbf"), "BOM utf-8-sig attendu en tête de fichier"

        lignes = brut.decode("utf-8-sig").splitlines()
        en_tetes = lignes[0].split(";")
        assert en_tetes[0] == "id"
        assert "complétude onboarding" in en_tetes
        attendu = len(service_riche.registry.list_nodes(projet.id))
        assert len([ligne for ligne in lignes[1:] if ligne]) == attendu


# --- Slugification du nom de fichier ------------------------------------------------------


class TestSlug:
    def test_accents_translitteres_et_slash_retire(
        self, service: SupplyScoreService, tmp_path: Path
    ) -> None:
        service.registry.save_project(
            Project(id="p-slug", name="Été à São Paulo / test", owner_node_id="n0")
        )

        path = ExportService(service).export_project("p-slug", dest_dir=tmp_path / "exports")

        assert path.name == f"SupplyScore_Ete-a-Sao-Paulo-test_{_STAMP}.xlsx"
        assert path.exists()
        # Le classeur d'un projet sans nœud reste valide (feuilles avec en-têtes).
        wb = load_workbook(path)
        assert wb.sheetnames == _FEUILLES


# --- Erreurs et destination ---------------------------------------------------------------


class TestErreursEtDestination:
    def test_projet_inconnu_value_error(self, service: SupplyScoreService, tmp_path: Path) -> None:
        with pytest.raises(ValueError, match="inconnu"):
            ExportService(service).export_project("fantome", dest_dir=tmp_path / "exports")

    def test_format_inconnu_value_error(
        self, exports: ExportService, projet: Project, tmp_path: Path
    ) -> None:
        with pytest.raises(ValueError, match="format"):
            exports.export_project(
                projet.id,
                fmt="pdf",  # type: ignore[arg-type]
                dest_dir=tmp_path / "exports",
            )

    def test_dest_dir_cree_au_besoin(
        self, exports: ExportService, projet: Project, tmp_path: Path
    ) -> None:
        dest = tmp_path / "exports" / "sous" / "dossier"
        assert not dest.exists()

        path = exports.export_project(projet.id, dest_dir=dest)

        assert dest.is_dir()
        assert path.parent == dest
