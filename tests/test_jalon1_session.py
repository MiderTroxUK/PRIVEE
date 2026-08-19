"""Critere de sortie du JALON 1 : une session de serious game complete est jouable.

Parcours : creer un projet en mode jeu -> onboarder un noeud au wizard (4 sections)
-> jouer des semaines -> remplir l'hebdo (AHP + evenement + decision, operateurs
identifies) -> suivre les scores -> exporter -> sauvegarder -> redemarrer sans perte.
"""

import zipfile

import pytest
from openpyxl import load_workbook

from supplyscore.core.clock import FixedClock
from supplyscore.data.backup import ServiceSauvegarde, restore_backup
from supplyscore.services import SupplyScoreService
from supplyscore.services.decisions import DecisionService
from supplyscore.services.events import EventEngine
from supplyscore.services.exports import ExportService
from supplyscore.services.onboarding import OnboardingService
from supplyscore.services.weekly import CycleHebdomadaire, StatutHebdo

T0 = 1_750_000_000.0


@pytest.fixture
def service(tmp_path):
    svc = SupplyScoreService(db_dir=tmp_path / "store", clock=FixedClock(T0))
    yield svc
    svc.close()


def test_session_serious_game_complete(service, tmp_path):
    # 1. Projet de base (la demo sert de chaine existante) + mode jeu.
    project = service.seed_demo(n_ranks=2, seed=21)
    service.set_clock_mode(project.id, "game")

    # 2. Onboarding d'un NOUVEAU fournisseur au wizard, rattache a un noeud existant.
    onboarding = OnboardingService(service)
    target = next(n for n in service.repo.nodes() if n.rank == 1)
    node_id = onboarding.start_draft(project.id, "Atelier Pilote", "Workshop")

    r1 = onboarding.save_section(
        node_id,
        1,
        {
            "name": "Atelier Pilote",
            "label": "Workshop",
            "location": "Toulouse",
            "tag_names": ["Pilote"],
            "connections": [{"target_id": target.id, "gamma": 0.6, "beta": 0.7, "kind": "nominal"}],
        },
        operator_id="animateur",
    )
    assert r1.ok, r1.errors

    week = 7 * 24 * 3600.0
    r2 = onboarding.save_section(
        node_id,
        2,
        {
            "cdc": {"budget_total": 50_000.0, "currency": "EUR"},
            "milestones": [
                {"name": "Proto", "kind": "proto", "start_ts": T0, "deadline_ts": T0 + 3 * week},
                {
                    "name": "Livraison",
                    "kind": "livraison",
                    "start_ts": T0 + 3 * week,
                    "deadline_ts": T0 + 8 * week,
                },
            ],
        },
        operator_id="animateur",
    )
    assert r2.ok, r2.errors

    r3 = onboarding.save_section(
        node_id,
        3,
        {"time.lead_time_h": 120.0, "time.lead_time_std_h": 24.0, "oee.availability": 0.92},
        operator_id="animateur",
    )
    assert r3.ok, r3.errors

    r4 = onboarding.save_section(
        node_id,
        4,
        {
            "comparisons": {"0-1": 1.0, "0-2": 1.0, "0-3": 1.0, "1-2": 1.0, "1-3": 1.0, "2-3": 1.0},
            "scores": [5.0, 5.0, 5.0, 5.0],
        },
        operator_id="joueur_atelier",
    )
    assert r4.ok, r4.errors
    rc = onboarding.complete(node_id, operator_id="animateur")
    assert rc.ok, rc.errors
    assert service.registry.get_node(node_id).onboarding_state == "complete"

    # 3. Tour de jeu : avancer 2 semaines - les statuts hebdo basculent.
    service.advance_week(project.id)
    service.advance_week(project.id)
    cycle = CycleHebdomadaire(service)
    etat = cycle.statut_noeud(node_id)
    assert etat.statut == StatutHebdo.EN_RETARD
    assert etat.semaines_de_retard == 2

    # 4. Hebdo du joueur : nouvel AHP + evenement calibre + decision tracee.
    assessment = SupplyScoreService.build_assessment(
        node_id=node_id,
        project_id=project.id,
        operator_id="joueur_atelier",
        comparisons={(0, 1): 1.0, (0, 2): 1.0, (0, 3): 1.0, (1, 2): 1.0, (1, 3): 1.0, (2, 3): 1.0},
        criteria_scores=[7.0, 7.0, 7.0, 7.0],
    )
    service.submit_assessment(assessment)
    assert cycle.statut_noeud(node_id).statut == StatutHebdo.A_JOUR

    engine = EventEngine(service)
    event = engine.apply(
        node_id,
        "retard_fournisseur",
        {"retard_h": 48.0},
        operator_id="joueur_atelier",
        notes="fournisseur matière en retard",
    )
    assert event.impacts, "l'événement doit impacter les KPIs"

    decision = DecisionService(service).record(
        node_id, "Doubler le stock de sécurité.", operator_id="joueur_atelier"
    )
    assert decision.scores["ur"] is not None

    # 5. Les scores vivent : le noeud a un etat complet.
    states = service.evaluate_all(persist=True)
    state = states[node_id]
    assert state.ud is not None and state.ur is not None and state.adequation is not None

    # 6. Export complet pour l'analyse post-jeu.
    export_path = ExportService(service).export_project(
        project.id, fmt="xlsx", dest_dir=tmp_path / "exports"
    )
    wb = load_workbook(export_path, read_only=True)
    assert {"noeuds", "evenements", "decisions", "evaluations"} <= set(wb.sheetnames)
    wb.close()

    # 7. Sauvegarde zip + restauration sur un repertoire neuf : aucune perte.
    zip_path = ServiceSauvegarde(service.db_dir, clock=service.clock).backup_all()
    assert zipfile.is_zipfile(zip_path)

    restored_dir = tmp_path / "restored"
    restore_backup(zip_path, restored_dir)
    reopened = SupplyScoreService(db_dir=restored_dir, clock=FixedClock(T0))
    try:
        reopened.load_graph_from_registry()
        node = reopened.registry.get_node(node_id)
        assert node is not None and node.name == "Atelier Pilote"
        decisions = reopened.client_db(node_id).list_decisions(node_id)
        assert len(decisions) == 1
        assert decisions[0]["description"] == "Doubler le stock de sécurité."
    finally:
        reopened.close()
